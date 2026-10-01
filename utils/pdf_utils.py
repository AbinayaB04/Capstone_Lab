import pymupdf
import io
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

def extract_and_chunk_pdf(file_bytes, chunk_size=1000, overlap=100, filename=""):
    """
    Extracts text from a document binary string/bytes and splits it into overlapping chunks per page.
    Returns a list of tuples: [(chunk_text, source_identifier)]
    """
    filename = filename.lower()
    chunks = []
    
    if filename.endswith(".ppt"):
        print(f"❌ Error: Legacy .ppt files are not supported. Please convert '{filename}' to .pptx or .pdf")
        return []
        
    try:
        if filename.endswith(".pptx"):
            if not Presentation:
                print("❌ Error: python-pptx is not installed. Cannot parse PPTX.")
                return []
            prs = Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(prs.slides):
                slide_text = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text += shape.text + "\n"
                
                start = 0
                while start < len(slide_text):
                    end = start + chunk_size
                    chunk = slide_text[start:end]
                    if chunk.strip():
                        chunks.append((chunk.strip(), f"Slide {i+1}"))
                    start += chunk_size - overlap
        else:
            pdf_document = pymupdf.open(stream=file_bytes, filetype="pdf")
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                page_text = page.get_text() + "\n"
                
                start = 0
                while start < len(page_text):
                    end = start + chunk_size
                    chunk = page_text[start:end]
                    if chunk.strip():
                        chunks.append((chunk.strip(), f"Page {page_num+1}"))
                    start += chunk_size - overlap
                    
        return chunks
    except Exception as e:
        print(f"Error parsing document: {e}")
        return []
